import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from backend.models.ad_revenues import AdRevenue
from backend.services.ad_revenues.features.export_ad_revenues import ExportAdRevenues


async def add_ad_revenue(db, owner_id, **overrides):
    row = AdRevenue(
        owner_id=owner_id,
        type=overrides.pop("type", "income"),
        amount=overrides.pop("amount", Decimal("100")),
        currency=overrides.pop("currency", "RUB"),
        revenue_date=overrides.pop("revenue_date", date(2026, 5, 15)),
        buyer=overrides.pop("buyer", "Acme"),
        **overrides,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return row


@pytest.mark.asyncio
async def test_xlsx_contains_section_sheet(db, test_user):
    await add_ad_revenue(db, test_user.id, type="income", amount=Decimal("250"))

    blob = await ExportAdRevenues(db).execute(
        owner_id=test_user.id,
        data_types={"general_income"},
        scope="filtered",
        export_format="xlsx",
    )

    wb = load_workbook(io.BytesIO(blob))
    assert "Общие доходы" in wb.sheetnames
    sheet = wb["Общие доходы"]
    # шапка + одна строка данных
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:6] == ("Дата", "Контрагент", "Сумма", "Валюта", "Канал", "Заметка")
    assert any(cell == "Acme" for cell in rows[1])


@pytest.mark.asyncio
async def test_xlsx_multiple_sections(db, test_user):
    await add_ad_revenue(db, test_user.id, type="income")
    await add_ad_revenue(db, test_user.id, type="expense", amount=Decimal("50"))

    blob = await ExportAdRevenues(db).execute(
        owner_id=test_user.id,
        data_types={"general_income", "general_expense"},
        scope="filtered",
        export_format="xlsx",
    )

    wb = load_workbook(io.BytesIO(blob))
    assert "Общие доходы" in wb.sheetnames
    assert "Общие расходы" in wb.sheetnames


@pytest.mark.asyncio
async def test_csv_contains_section_titles_and_data(db, test_user):
    await add_ad_revenue(db, test_user.id, type="income", buyer="Acme")

    blob = await ExportAdRevenues(db).execute(
        owner_id=test_user.id,
        data_types={"general_income"},
        scope="filtered",
        export_format="csv",
    )

    text = blob.decode("utf-8").lstrip("﻿")
    reader = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert ["Общие доходы"] in reader
    assert any("Acme" in cell for row in reader for cell in row)


@pytest.mark.asyncio
async def test_scope_all_ignores_filters(db, test_user):
    """scope='all' должен выгрузить всё, игнорируя date_from."""
    await add_ad_revenue(db, test_user.id, revenue_date=date(2020, 1, 1))

    blob = await ExportAdRevenues(db).execute(
        owner_id=test_user.id,
        data_types={"general_income"},
        scope="all",
        export_format="csv",
        date_from=date(2026, 1, 1),
    )

    text = blob.decode("utf-8").lstrip("﻿")
    assert "Acme" in text
