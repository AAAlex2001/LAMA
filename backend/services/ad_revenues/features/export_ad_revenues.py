import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import List, Literal, Optional, Set

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from backend.schemas.ad_revenues.ad_revenue import AdRevenueResponse
from backend.schemas.ad_revenues.enums import AdRevenueType
from backend.services.ad_revenues.features.list_ad_revenues import ListAdRevenues


ExportFormat = Literal["xlsx", "csv"]
ExportScope = Literal["filtered", "all"]
DataType = Literal["general_income", "general_expense", "ads_income", "ads_expense"]

GENERAL_COLUMNS = ["Дата", "Контрагент", "Сумма", "Валюта", "Канал", "Заметка"]
ADS_COLUMNS = GENERAL_COLUMNS + ["Просмотры", "Пересылки", "Реакции", "Комментарии", "Клики", "Ссылка"]

SECTION_TITLES = {
    "general_income": "Общие доходы",
    "general_expense": "Общие расходы",
    "ads_income": "Доходы рекламы",
    "ads_expense": "Расходы рекламы",
}


@dataclass(frozen=True)
class ExportSection:
    key: DataType
    title: str
    columns: List[str]
    rows: List[List[object]]


class ExportAdRevenues:
    """Экспорт рекламных записей в xlsx или csv с разделением по секциям."""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    async def execute(
        self,
        owner_id: int,
        data_types: Set[str],
        scope: ExportScope,
        export_format: ExportFormat,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        channel_id: Optional[int] = None,
        bot_id: Optional[int] = None,
        currency: Optional[str] = None,
    ) -> bytes:
        applied_from = date_from if scope == "filtered" else None
        applied_to = date_to if scope == "filtered" else None
        applied_channel = channel_id if scope == "filtered" else None
        applied_bot = bot_id if scope == "filtered" else None
        applied_currency = currency if scope == "filtered" else None

        sections: List[ExportSection] = []
        for key in ("general_income", "general_expense", "ads_income", "ads_expense"):
            if key not in data_types:
                continue
            section = await self.build_section(
                key=key,  # type: ignore[arg-type]
                owner_id=owner_id,
                date_from=applied_from,
                date_to=applied_to,
                channel_id=applied_channel,
                bot_id=applied_bot,
                currency=applied_currency,
            )
            sections.append(section)

        if export_format == "xlsx":
            return self.render_xlsx(sections)
        return self.render_csv(sections)

    async def build_section(
        self,
        key: DataType,
        owner_id: int,
        date_from: Optional[date],
        date_to: Optional[date],
        channel_id: Optional[int],
        bot_id: Optional[int],
        currency: Optional[str],
    ) -> ExportSection:
        type_filter = AdRevenueType.INCOME if key.endswith("income") else AdRevenueType.EXPENSE
        items, _ = await ListAdRevenues(self.db).execute(
            owner_id=owner_id,
            type_=type_filter,
            channel_id=channel_id,
            bot_id=bot_id,
            date_from=date_from,
            date_to=date_to,
            currency=currency,
            limit=10000,
            offset=0,
        )
        is_ads_section = key.startswith("ads")
        if is_ads_section:
            items = [r for r in items if r.publication_id is not None]
            columns = ADS_COLUMNS
            rows = [ads_row(r) for r in items]
        else:
            columns = GENERAL_COLUMNS
            rows = [general_row(r) for r in items]
        return ExportSection(key=key, title=SECTION_TITLES[key], columns=columns, rows=rows)

    def render_xlsx(self, sections: List[ExportSection]) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)
        for section in sections:
            ws = wb.create_sheet(title=truncate_sheet_title(section.title))
            ws.append(section.columns)
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
            for row in section.rows:
                ws.append(row)
            for column_idx in range(1, len(section.columns) + 1):
                ws.column_dimensions[chr(64 + column_idx)].width = 18
        if not wb.sheetnames:
            wb.create_sheet(title="Empty")
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def render_csv(self, sections: List[ExportSection]) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
        for index, section in enumerate(sections):
            if index > 0:
                writer.writerow([])
            writer.writerow([section.title])
            writer.writerow(section.columns)
            for row in section.rows:
                writer.writerow(row)
        return ("﻿" + buffer.getvalue()).encode("utf-8")


def general_row(item: AdRevenueResponse) -> List[object]:
    return [
        format_date(item.revenue_date),
        item.buyer or "",
        format_amount(item.amount),
        item.currency,
        placement_label(item),
        item.note or "",
    ]


def ads_row(item: AdRevenueResponse) -> List[object]:
    return [
        format_date(item.revenue_date),
        item.buyer or "",
        format_amount(item.amount),
        item.currency,
        placement_label(item),
        item.note or "",
        item.views_count,
        item.forwards_count,
        item.reactions_count,
        item.comments_count,
        item.clicks_count,
        item.post_link or "",
    ]


def placement_label(item: AdRevenueResponse) -> str:
    if item.placements:
        names = [p.title or (f"@{p.username}" if p.username else "") for p in item.placements]
        return ", ".join(name for name in names if name)
    if item.channel_username:
        return f"@{item.channel_username.lstrip('@')}"
    return ""


def format_date(value: date) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat()


def format_amount(value: Decimal) -> float:
    if value is None:
        return 0.0
    return float(value)


def truncate_sheet_title(title: str) -> str:
    return title[:31]
